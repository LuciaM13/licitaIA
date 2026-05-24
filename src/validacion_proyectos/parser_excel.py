"""Parser de los Excel BC3 individuales de EMASESA hacia ProyectoReal.

Detecta filas hoja "Total partida X" (cualquier código), recupera la
descripción real desde la fila de cabecera de la partida (no la fila Total),
y agrupa partidas por **capítulo top corriente** (el cap top en col A más
reciente al recorrer en orden), no por primer segmento del código.

Las columnas siguen la tabla de data/proyectos_individuales/README.md (A-K).
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from openpyxl import load_workbook

from src.validacion_proyectos.tipos import PartidaReal, ProyectoReal

logger = logging.getLogger(__name__)


# Acepta cualquier código tras "Total partida" (incluye GOB.05.2025.001, S01, 1.3.05, etc.)
_RE_TOTAL_PARTIDA = re.compile(r"^\s*Total\s+partida\s+(\S+?)\s*$", re.IGNORECASE)
# Cap top: col A es un código numérico de 1-2 dígitos sin puntos
_RE_CAP_TOP = re.compile(r"^\d{1,2}$")


def _parsear_decimal_es(valor: object) -> float:
    if valor is None:
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    s = str(valor).strip()
    if not s:
        return 0.0
    s = s.replace(".", "").replace(",", ".") if ("," in s and s.count(",") == 1) else s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        logger.warning("[PARSER] No se pudo parsear decimal: %r -> 0.0", valor)
        return 0.0


def parsear_proyecto(path: Path) -> ProyectoReal:
    logger.info("parsear_proyecto INICIO path=%s", path)
    path = Path(path)
    wb = load_workbook(path, data_only=True)  # NO read_only: necesitamos lookback
    ws = wb["Hoja1"] if "Hoja1" in wb.sheetnames else wb.active

    # Pasada 1: índice cabecera_por_codigo (cod col A -> descripción col C)
    # y secuencia ordenada de filas.
    rows = list(ws.iter_rows(values_only=True))
    cabecera_por_codigo: dict[str, str] = {}
    cap_top_por_fila: dict[int, str] = {}
    cap_top_actual = ""
    for idx, fila in enumerate(rows, start=1):
        if not fila or len(fila) < 3:
            cap_top_por_fila[idx] = cap_top_actual
            continue
        a, c = fila[0], fila[2]
        a_s = str(a).strip() if a is not None else ""
        c_s = c.strip() if isinstance(c, str) else ""
        # Detección cap top: col A coincide con regex \d{1,2} y col C es no-Total
        if _RE_CAP_TOP.match(a_s) and c_s and not c_s.lower().startswith("total"):
            cap_top_actual = a_s.zfill(2)
        cap_top_por_fila[idx] = cap_top_actual
        # Index cabecera (col A no vacío, col C no-Total)
        if a_s and a_s != "." and c_s and not c_s.lower().startswith("total"):
            cabecera_por_codigo[a_s] = c_s

    # Pasada 2: extraer partidas hoja
    partidas: list[PartidaReal] = []
    for idx, fila in enumerate(rows, start=1):
        if idx <= 2:
            continue
        if not fila or len(fila) < 11:
            continue
        col_c = fila[2]
        if not isinstance(col_c, str):
            continue
        m = _RE_TOTAL_PARTIDA.match(col_c)
        if not m:
            continue
        codigo = m.group(1).strip()
        descripcion_real = cabecera_por_codigo.get(codigo, col_c.strip())
        capitulo = cap_top_por_fila.get(idx, "")
        if not capitulo:
            # Fallback: primer segmento del código (legacy comportamiento)
            capitulo = codigo.split(".")[0].zfill(2)

        medicion = _parsear_decimal_es(fila[8])
        precio_unitario = _parsear_decimal_es(fila[9])
        importe = _parsear_decimal_es(fila[10])

        partidas.append(PartidaReal(
            codigo=codigo,
            capitulo=capitulo,
            descripcion=descripcion_real,
            medicion=medicion,
            precio_unitario=precio_unitario,
            importe=importe,
        ))

    wb.close()

    capitulos_real: dict[str, float] = {}
    for p in partidas:
        capitulos_real[p.capitulo] = capitulos_real.get(p.capitulo, 0.0) + p.importe

    pem_total = sum(p.importe for p in partidas)
    expediente = path.stem.split(" ")[0] if path.stem else path.stem
    nombre = path.stem

    proyecto = ProyectoReal(
        expediente=expediente,
        nombre=nombre,
        fuente_xlsx=path,
        pem_total=pem_total,
        capitulos_real=capitulos_real,
        partidas=partidas,
    )
    logger.info("parsear_proyecto FIN expediente=%s partidas=%d pem_total=%.2f caps=%s",
                expediente, len(partidas), pem_total, sorted(capitulos_real.keys()))
    return proyecto
